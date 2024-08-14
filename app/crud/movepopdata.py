import mysql.connector
from app.db.connect import get_db_connection  # Ensure this function returns a MySQL connection

def read_keywords_from_excel(file_path):
    import openpyxl
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    keywords = []
    for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0] is not None:
            keywords.append(row[0])
    return keywords

def insert_record(table_name: str, **kwargs):
    connection = get_db_connection()
    try:
        with connection.cursor() as cursor:
            columns = ', '.join(kwargs.keys())
            placeholders = ', '.join(['%s'] * len(kwargs))
            values = tuple(kwargs.values())
            sql = f"INSERT INTO {table_name} ({columns}) VALUES ({placeholders})"
            cursor.execute(sql, values)
        connection.commit()
    except mysql.connector.IntegrityError as e:
        if e.errno == 1062:  # Duplicate entry error code
            print(f"Duplicate entry found for {kwargs}. Skipping...")
        else:
            print(f"Error inserting data: {e}")
            connection.rollback()
            raise e
    except Exception as e:
        print(f"Error inserting data: {e}")
        connection.rollback()
        raise e
    finally:
        connection.close()
